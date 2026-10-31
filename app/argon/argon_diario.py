"""Ciclo diario de la fase 1: lee la Excel, compara con la lectura anterior y prepara el correo de avisos.

Uso:  python argon_diario.py ARGON.xlsx carpeta_estado [--hoy 2026-10-30]
  carpeta_estado guarda la última lectura (ultima_lectura.json) para comparar la próxima vez.
Salida: cambios.json (lo que iría a Odoo), aviso.html (correo al equipo) y recordatorios.json
(un correo por destinatario con los vencimientos ADR/ITV que toca recordar hoy).
"""
from __future__ import annotations

import datetime as dt
import html
import json
import sys
from collections import defaultdict
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
import argon_lector as L  # noqa: E402
import openpyxl  # noqa: E402
import argon_facturacion as F  # noqa: E402

VENTANA_DIAS = 45        # se comparan los movimientos desde hoy-45 hasta el final (planificados incluidos)
DIAS_SIN_ANOTAR = 3


def comparar(antes: list[dict], ahora: list[dict]) -> dict:
    a = {m["id_evento"]: m for m in antes}
    b = {m["id_evento"]: m for m in ahora}
    campos = ("hora", "carga", "bloqueado", "texto", "comentario", "chasis", "iniciales")
    nuevos = [b[k] for k in b if k not in a]
    borrados = [a[k] for k in a if k not in b]
    cambiados = [dict(b[k], antes={c: a[k].get(c) for c in campos if a[k].get(c) != b[k].get(c)})
                 for k in b if k in a and any(a[k].get(c) != b[k].get(c) for c in campos)]
    return {"nuevos": nuevos, "cambiados": cambiados, "borrados": borrados}


def avisos(conts, estados, movs_hoy, hoy: dt.date) -> dict:
    por_cont = defaultdict(list)
    for e in estados:
        por_cont[e.contenedor].append(e)
    sin_anotar, bloqueados = [], []
    for k in conts:
        lst = sorted((e for e in por_cont.get(k.matricula, []) if e.fecha <= hoy.isoformat()), key=lambda e: e.fecha)
        if not lst or lst[-1].tipo == "FUERA_SERVICIO":
            continue
        ult = lst[-1]
        dias = (hoy - dt.date.fromisoformat(ult.fecha)).days
        if dias >= DIAS_SIN_ANOTAR:
            sin_anotar.append((k, ult, dias))
        if ult.bloqueado and ult.fecha == hoy.isoformat():
            bloqueados.append((k, ult))
    desde = (hoy - dt.timedelta(days=7)).isoformat()
    revisar = [e for e in estados if e.fecha >= desde and e.avisos]
    return {"sin_anotar": sin_anotar, "bloqueados": bloqueados, "revisar": revisar}


DIAS_RECORDATORIO = {30, 15, 7, 1, 0}


def _sumar_meses(f: dt.date, meses: int) -> dt.date:
    m = f.month - 1 + meses
    y, m = f.year + m // 12, m % 12 + 1
    import calendar
    return dt.date(y, m, min(f.day, calendar.monthrange(y, m)[1]))


def vencimientos(ruta, hoy: dt.date) -> list[dict]:
    wb = openpyxl.load_workbook(ruta, data_only=False)
    if "VENCIMIENTOS" not in wb.sheetnames:
        return []
    res = []
    for fila in wb["VENCIMIENTOS"].iter_rows(min_row=4, values_only=True):
        el, ident, tipo, ultima, period, exacta = fila[:6]
        if not el:
            continue
        d = lambda v: v.date() if isinstance(v, dt.datetime) else v if isinstance(v, dt.date) else None  # noqa: E731
        vence = d(exacta) or (_sumar_meses(d(ultima), int(period)) if d(ultima) and isinstance(period, (int, float)) and period else None)
        if not vence:
            continue
        dias = (vence - hoy).days
        aviso = int(fila[9]) if isinstance(fila[9], (int, float)) else 30
        estado = "VENCIDO" if dias < 0 else "AVISAR" if dias <= aviso else "OK"
        correos = [c.strip() for c in str(fila[10] or "").replace(",", ";").split(";") if "@" in c]
        recordar = dias in (DIAS_RECORDATORIO | {aviso}) or (dias < 0 and hoy.weekday() == 0)
        res.append({"elemento": el, "id": str(ident), "tipo": tipo, "vence": vence.isoformat(), "dias": dias,
                    "estado": estado, "correos": correos, "recordar": recordar})
    return res


def recordatorios(venc: list[dict], hoy: dt.date) -> list[dict]:
    """Un correo por destinatario con lo que toca recordarle hoy (n8n los envía tal cual)."""
    por = defaultdict(list)
    for v in venc:
        if v["recordar"]:
            for c in v["correos"]:
                por[c].append(v)
    salida = []
    for para, items in por.items():
        filas = "".join(f"<tr><td>{html.escape(i['elemento'])} {html.escape(i['id'])}</td><td>{html.escape(i['tipo'])}</td>"
                        f"<td>{dt.date.fromisoformat(i['vence']):%d/%m/%Y}</td><td>"
                        f"{'vencido hace ' + str(-i['dias']) + ' días' if i['dias'] < 0 else 'hoy' if i['dias'] == 0 else 'faltan ' + str(i['dias']) + ' días'}"
                        f"</td></tr>" for i in sorted(items, key=lambda x: x["dias"]))
        salida.append({"para": para, "asunto": f"ARGON · vencimientos ADR/ITV ({hoy:%d/%m/%Y})",
                       "html": f"<p>Recordatorio de vencimientos:</p><table border='1' cellpadding='4' "
                               f"style='border-collapse:collapse;font:10pt Calibri,Arial'><tr><th>Elemento</th><th>Tipo</th>"
                               f"<th>Vence</th><th></th></tr>{filas}</table>"})
    return salida


def html_aviso(hoy, cambios, av) -> str:
    e = html.escape
    f = lambda d: dt.date.fromisoformat(d).strftime("%d/%m")  # noqa: E731
    td = 'style="padding:3px 8px;border-bottom:1px solid #ddd"'
    th = 'style="padding:4px 8px;background:#1F4E79;color:#fff;text-align:left"'

    def tabla(cab, filas):
        if not filas:
            return "<p style='color:#555'>Nada.</p>"
        h = "".join(f"<th {th}>{e(c)}</th>" for c in cab)
        b = "".join("<tr>" + "".join(f"<td {td}>{e(str(v or ''))}</td>" for v in fila) + "</tr>" for fila in filas)
        return f"<table style='border-collapse:collapse;font:10pt Calibri,Arial'><tr>{h}</tr>{b}</table>"

    n, c, b = len(cambios["nuevos"]), len(cambios["cambiados"]), len(cambios["borrados"])
    partes = [f"<p style='font:11pt Calibri,Arial'><b>ARGON · lectura del {hoy:%d/%m/%Y}</b><br>"
              f"Para Odoo: {n} movimientos nuevos, {c} modificados y {b} que ya no están en la Excel.</p>"]
    partes.append("<h3 style='font:bold 11pt Calibri,Arial'>Bloqueados hoy</h3>")
    partes.append(tabla(["Contenedor", "Código", "Motivo"], [(k.matricula, k.codigo_corto, u.texto) for k, u in av["bloqueados"]]))
    partes.append(f"<h3 style='font:bold 11pt Calibri,Arial'>Sin anotar desde hace {DIAS_SIN_ANOTAR} días o más</h3>")
    partes.append(tabla(["Contenedor", "Código", "Última anotación", "Fecha", "Días"],
                        [(k.matricula, k.codigo_corto, u.texto, f(u.fecha), d) for k, u, d in sorted(av["sin_anotar"], key=lambda x: -x[2])]))
    partes.append("<h3 style='font:bold 11pt Calibri,Arial'>Para revisar en la Excel (últimos 7 días)</h3>")
    partes.append(tabla(["Fecha", "Contenedor", "Celda", "Texto", "Aviso"],
                        [(f(x.fecha), x.contenedor, x.celda, x.texto, "; ".join(x.avisos)) for x in av["revisar"]]))
    partes.append("<h3 style='font:bold 11pt Calibri,Arial'>Salidas en tren sin nº de expedición (últimos 7 días)</h3>")
    partes.append(tabla(["Fecha", "Contenedor", "Celda", "Texto"],
                        [(f(x.fecha), x.contenedor, x.celda, x.texto) for x in av.get("sl_sin_numero", [])]))
    if av.get("reservas_sin_salida") is not None:
        partes.append("<h3 style='font:bold 11pt Calibri,Arial'>Reservas de tren sin salida anotada (±2 días)</h3>")
        partes.append(tabla(["Expedición", "Proveedor", "Salida", "Origen", "Contenedor"],
                            [(r.get("Nº expedición"), r.get("Proveedor"), r["_fecha"].strftime("%d/%m") if r.get("_fecha") else "",
                              r.get("Origen"), r.get("Contenedor")) for r in av["reservas_sin_salida"]]))
    venc = [v for v in av.get("vencimientos", []) if v["estado"] != "OK"]
    partes.append("<h3 style='font:bold 11pt Calibri,Arial'>Vencimientos ADR / ITV</h3>")
    partes.append(tabla(["Elemento", "Tipo", "Vence", "Días", "Estado"],
                        [(f"{v['elemento']} {v['id']}", v["tipo"], f(v["vence"]), v["dias"], v["estado"])
                         for v in sorted(venc, key=lambda x: x["dias"])]))
    partes.append("<p style='font:9pt Calibri,Arial;color:#777'>Si es un sitio nuevo, añadidlo en la hoja UBICACIONES; "
                  "la próxima lectura ya lo entenderá.</p>")
    return "\n".join(partes)


def main(argv):
    if len(argv) < 3:
        print(__doc__)
        return 1
    ruta, carpeta = Path(argv[1]), Path(argv[2])
    hoy = dt.date.fromisoformat(argv[argv.index("--hoy") + 1]) if "--hoy" in argv else dt.date.today()
    carpeta.mkdir(parents=True, exist_ok=True)
    conts, estados = L.leer(ruta)
    movs = [m for m in L.movimientos(estados) if m["fecha"] >= (hoy - dt.timedelta(days=VENTANA_DIAS)).isoformat()]
    previa = carpeta / "ultima_lectura.json"
    antes = json.loads(previa.read_text(encoding="utf-8")) if previa.exists() else []
    desde = (hoy - dt.timedelta(days=VENTANA_DIAS)).isoformat()
    antes = [m for m in antes if m["fecha"] >= desde]
    cambios = comparar(antes, movs)
    (carpeta / "cambios.json").write_text(json.dumps(cambios, ensure_ascii=False, indent=1), encoding="utf-8")
    av = avisos(conts, estados, movs, hoy)
    desde7 = (hoy - dt.timedelta(days=7)).isoformat()
    av["sl_sin_numero"] = [e for e in estados if e.tipo == "SALIDA_TREN" and not e.referencia and desde7 <= e.fecha <= hoy.isoformat()]
    tablas = F.leer_tablas(ruta)
    if tablas["reservas"]:
        servs = F.construir(L.movimientos(estados), tablas["tipos"])
        cruce = F.cruzar_reservas(servs, tablas["reservas"])
        av["reservas_sin_salida"] = [r for r in cruce["reservas_sin_salida"] if r.get("_fecha") and r["_fecha"] <= hoy
                                     and L.norm(r.get("Estado") or "") != "CANCELADA"]
        con_reserva = {(s.contenedor, s.fecha) for s in servs if s.reserva}
        av["sl_sin_numero"] = [e for e in av["sl_sin_numero"] if (e.contenedor, e.fecha) not in con_reserva]
    av["vencimientos"] = vencimientos(ruta, hoy)
    (carpeta / "aviso.html").write_text(html_aviso(hoy, cambios, av), encoding="utf-8")
    (carpeta / "recordatorios.json").write_text(json.dumps(recordatorios(av["vencimientos"], hoy), ensure_ascii=False, indent=1),
                                                encoding="utf-8")
    previa.write_text(json.dumps(movs, ensure_ascii=False), encoding="utf-8")
    print(json.dumps({k: len(v) for k, v in cambios.items()}))
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
